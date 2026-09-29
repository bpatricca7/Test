"""End-to-end tests: FastAPI app + real openai SDK + simulator + Jupyter sandbox + MCP."""

from __future__ import annotations

import io
import json
import shutil
import sys
from pathlib import Path

import pytest
from fastapi.testclient import TestClient

ROOT = Path(__file__).resolve().parents[2]


def types(events):
    return [e["type"] for e in events]


def test_demo_turn_with_code_interpreter(make_app, chat):
    with TestClient(make_app()) as c:
        events = chat(c, {"message": "Make a chart of planned vs actual hours", "model": "gpt-5.6-terra"})
        t = types(events)
        assert t[0] == "meta" and "done" in t and t.index("done") > t.index("usage")
        assert "reasoning_delta" in t and "tool_input_delta" in t and "text_delta" in t
        tool_end = next(e for e in events if e["type"] == "tool_end")
        assert tool_end["status"] == "done", tool_end
        assert tool_end["output"]["images"], "matplotlib chart should be captured"
        assert [f["name"] for f in tool_end["output"]["files"]] == ["labor_hours_demo.xlsx"]
        usage = next(e for e in events if e["type"] == "usage")["usage"]
        assert usage["calls"] == 2 and usage["cost_usd"] > 0 and usage["reasoning_tokens"] > 0
        assert any(e["type"] == "title" for e in events)

        cid = events[0]["conversation_id"]
        conv = c.get(f"/api/conversations/{cid}").json()
        assistant = conv["messages"][-1]
        assert assistant["status"] == "complete"
        assert [p["type"] for p in assistant["parts"]] == ["reasoning", "tool", "reasoning", "text"]
        file_id = tool_end["output"]["files"][0]["id"]
        assert c.get(f"/api/files/{file_id}").content[:2] == b"PK"  # xlsx zip

        # Follow-up turn in the same conversation reuses the kernel state and history.
        events2 = chat(c, {"message": "thanks!", "conversation_id": cid})
        assert "done" in types(events2)
        usage_report = c.get("/api/usage").json()
        assert usage_report["totals"]["requests"] == 2
        assert {m["model"] for m in usage_report["by_model"]} >= {"gpt-5.6-terra"}


def test_office_upload_edit_and_print(make_app, chat):
    import docx

    doc = docx.Document()
    doc.add_paragraph("Program review draft")
    buf = io.BytesIO()
    doc.save(buf)
    with TestClient(make_app()) as c:
        up = c.post("/api/files", files=[("files", ("review.docx", buf.getvalue(), "application/octet-stream"))]).json()
        assert up[0]["kind"] == "word"
        events = chat(c, {"message": "Please edit this document", "attachments": [up[0]["id"]]})
        end = next(e for e in events if e["type"] == "tool_end")
        assert end["status"] == "done", end["output"]
        names = [f["name"] for f in end["output"]["files"]]
        assert "review_edited.docx" in names
        edited = next(f for f in end["output"]["files"] if f["name"] == "review_edited.docx")
        data = c.get(f"/api/files/{edited['id']}").content
        assert "Reviewed with Amentum AI" in "\n".join(p.text for p in docx.Document(io.BytesIO(data)).paragraphs)
        if shutil.which("soffice"):
            pdf = c.get(f"/api/files/{edited['id']}/pdf")
            assert pdf.status_code in (200, 422)
            if pdf.status_code == 200:
                assert pdf.content[:4] == b"%PDF"


def test_mcp_connector_end_to_end(make_app, chat):
    with TestClient(make_app()) as c:
        server = c.post("/api/connectors", json={
            "name": "Portfolio", "transport": "stdio", "command": sys.executable,
            "args": [str(ROOT / "mcp-servers" / "amentum-demo" / "server.py")],
            "description": "Program data", "env": {"SECRET_TOKEN": "abc"},
        }).json()
        assert server["status"] == "connected", server
        assert server["env"]["SECRET_TOKEN"] != "abc", "secrets must be masked"
        assert {t["name"] for t in server["tools"]} >= {"list_open_risks", "search_programs"}

        events = chat(c, {"message": "What are the top program risks?", "connectors": [server["id"]]})
        start = next(e for e in events if e["type"] == "tool_start")
        assert start["kind"] == "mcp" and start["server"] == "Portfolio" and start["name"] == "list_open_risks"
        end = next(e for e in events if e["type"] == "tool_end")
        assert end["status"] == "done" and "Groundwater" in json.dumps(end["output"])
        text = "".join(e["delta"] for e in events if e["type"] == "text_delta")
        assert "| program_id |" in text

        # Disabled connectors are not offered to the model.
        c.put(f"/api/connectors/{server['id']}", json={"enabled": False})
        events = chat(c, {"message": "What are the top program risks?"})
        assert not any(e["type"] == "tool_start" and e["kind"] == "mcp" for e in events)
        assert c.delete(f"/api/connectors/{server['id']}").json() == {"ok": True}


def _chat_models_file(tmp_path: Path) -> str:
    data = json.loads((ROOT / "config" / "models.json").read_text())
    for m in data["models"]:
        m["api"] = "chat"
    p = tmp_path / "models.chat.json"
    p.write_text(json.dumps(data))
    return str(p)


def test_gcc_high_versioned_chat_completions_path(make_app, chat, fake_llm_url, tmp_path):
    """AsyncAzureOpenAI (api-version style) + Chat Completions + tool calls."""
    app = make_app(LLM_PROVIDER="azure_gcc_high", ENFORCE_GOV_ENDPOINTS="false",
                   AZURE_OPENAI_ENDPOINT=f"{fake_llm_url}/az", AZURE_OPENAI_API_KEY="test-key",
                   AZURE_API_STYLE="versioned", MODELS_FILE=_chat_models_file(tmp_path))
    with TestClient(app) as c:
        assert c.get("/api/config").json()["provider_label"] == "Azure Government · GCC High"
        events = chat(c, {"message": "chart the hours please", "model": "gpt-5.6-sol", "effort": "high"})
        t = types(events)
        assert "reasoning_delta" in t, "reasoning_content deltas become thinking traces"
        end = next(e for e in events if e["type"] == "tool_end")
        assert end["status"] == "done"
        assert next(e for e in events if e["type"] == "done")["status"] == "complete"
        assert next(e for e in events if e["type"] == "usage")["usage"]["calls"] == 2


def test_gcc_high_v1_responses_path(make_app, chat, fake_llm_url):
    """OpenAI client against the Azure v1 route + Responses API reasoning summaries."""
    app = make_app(LLM_PROVIDER="azure_gcc_high", ENFORCE_GOV_ENDPOINTS="false",
                   AZURE_OPENAI_ENDPOINT=f"{fake_llm_url}/az", AZURE_OPENAI_API_KEY="test-key", AZURE_API_STYLE="v1")
    with TestClient(app) as c:
        events = chat(c, {"message": "hello there", "model": "gpt-5.6-luna", "code_interpreter": False})
        reasoning = "".join(e["delta"] for e in events if e["type"] == "reasoning_delta")
        assert "Understanding the request" in reasoning
        assert next(e for e in events if e["type"] == "done")["status"] == "complete"


def test_conversation_management(make_app, chat):
    with TestClient(make_app()) as c:
        events = chat(c, {"message": "hi", "code_interpreter": False})
        cid = events[0]["conversation_id"]
        assert c.patch(f"/api/conversations/{cid}", json={"title": "Renamed", "pinned": True}).json()["ok"]
        listed = c.get("/api/conversations").json()
        assert listed[0]["title"] == "Renamed" and listed[0]["pinned"] is True
        assert c.get("/api/conversations", params={"q": "zzz-nomatch"}).json() == []
        # Regenerate: re-run from the user message.
        user_msg = events[0]["user_message_id"]
        chat(c, {"message": "hi again", "conversation_id": cid, "edit_message_id": user_msg, "code_interpreter": False})
        msgs = c.get(f"/api/conversations/{cid}").json()["messages"]
        assert [m["role"] for m in msgs] == ["user", "assistant"] and msgs[0]["content"] == "hi again"
        assert c.delete(f"/api/conversations/{cid}").json()["ok"]
        assert c.get(f"/api/conversations/{cid}").status_code == 404


def test_diagnostics_probe(make_app):
    with TestClient(make_app()) as c:
        d = c.get("/api/diagnostics", params={"probe": True}).json()
        assert d["probe"]["ok"] is True
        assert d["sandbox"]["ok"] is True
