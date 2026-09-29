"""File storage and document text extraction."""

from __future__ import annotations

import mimetypes
import re
from pathlib import Path

from .config import get_settings
from .db import get_db, new_id

OFFICE_KINDS = {
    ".docx": "word", ".doc": "word", ".rtf": "word", ".odt": "word",
    ".xlsx": "excel", ".xlsm": "excel", ".xls": "excel", ".csv": "excel", ".tsv": "excel", ".ods": "excel",
    ".pptx": "powerpoint", ".ppt": "powerpoint", ".odp": "powerpoint",
    ".pdf": "pdf",
    ".png": "image", ".jpg": "image", ".jpeg": "image", ".gif": "image", ".webp": "image", ".svg": "image",
    ".txt": "text", ".md": "text", ".json": "code", ".xml": "code", ".py": "code", ".html": "code",
    ".zip": "archive",
}
EXTRA_MIME = {
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ".pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    ".md": "text/markdown",
    ".csv": "text/csv",
}
IMAGE_MIMES = {"image/png", "image/jpeg", "image/gif", "image/webp"}
CONVERTIBLE_TO_PDF = {".docx", ".doc", ".rtf", ".odt", ".xlsx", ".xls", ".ods", ".csv", ".pptx", ".ppt", ".odp", ".txt", ".md", ".html"}


def safe_name(name: str) -> str:
    name = Path(name.replace("\\", "/")).name
    name = re.sub(r"[^\w\-. ()\[\]&+,@]", "_", name).strip(" .")
    return name[:180] or "file"


def guess_mime(name: str) -> str:
    ext = Path(name).suffix.lower()
    return EXTRA_MIME.get(ext) or mimetypes.guess_type(name)[0] or "application/octet-stream"


def file_kind(name: str) -> str:
    return OFFICE_KINDS.get(Path(name).suffix.lower(), "file")


def storage_dir() -> Path:
    d = get_settings().data_path / "files"
    d.mkdir(parents=True, exist_ok=True)
    return d


def store_bytes(
    user: str,
    name: str,
    data: bytes,
    source: str = "upload",
    conversation_id: str | None = None,
    sandbox_path: str | None = None,
) -> dict:
    fid = new_id("file")
    name = safe_name(name)
    folder = storage_dir() / (conversation_id or "_pending")
    folder.mkdir(parents=True, exist_ok=True)
    path = folder / f"{fid}__{name}"
    path.write_bytes(data)
    return get_db().add_file(
        user=user,
        name=name,
        mime=guess_mime(name),
        size=len(data),
        path=str(path),
        source=source,
        conversation_id=conversation_id,
        sandbox_path=sandbox_path,
        file_id=fid,
    )


def public_file(f: dict) -> dict:
    ext = Path(f["name"]).suffix.lower()
    return {
        "id": f["id"],
        "name": f["name"],
        "mime": f["mime"],
        "size": f["size"],
        "kind": file_kind(f["name"]),
        "source": f["source"],
        "sandbox_path": f.get("sandbox_path"),
        "previewable": f["mime"] in IMAGE_MIMES or ext in CONVERTIBLE_TO_PDF or ext in (".pdf", ".svg"),
        "url": f"/api/files/{f['id']}",
    }


def extract_text(path: str | Path, limit: int | None = None) -> str:
    """Best-effort plain-text extraction used when the code interpreter is disabled."""
    path = Path(path)
    limit = limit or get_settings().max_inline_text_chars
    ext = path.suffix.lower()
    try:
        if ext == ".docx":
            import docx

            d = docx.Document(str(path))
            chunks = [p.text for p in d.paragraphs]
            for t in d.tables:
                for row in t.rows:
                    chunks.append(" | ".join(c.text.strip() for c in row.cells))
            text = "\n".join(chunks)
        elif ext in (".xlsx", ".xlsm"):
            import openpyxl

            wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
            out = []
            for ws in wb.worksheets:
                out.append(f"## Sheet: {ws.title}")
                for i, row in enumerate(ws.iter_rows(values_only=True)):
                    if i > 500:
                        out.append("… (truncated)")
                        break
                    out.append(",".join("" if v is None else str(v) for v in row))
            text = "\n".join(out)
        elif ext == ".pptx":
            from pptx import Presentation

            prs = Presentation(str(path))
            out = []
            for i, slide in enumerate(prs.slides, 1):
                out.append(f"## Slide {i}")
                for shape in slide.shapes:
                    if getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip():
                        out.append(shape.text_frame.text)
            text = "\n".join(out)
        elif ext == ".pdf":
            from pypdf import PdfReader

            reader = PdfReader(str(path))
            text = "\n".join((p.extract_text() or "") for p in reader.pages[:200])
        else:
            text = path.read_text(encoding="utf-8", errors="replace")
    except Exception as exc:  # pragma: no cover - depends on file contents
        return f"[Could not extract text from {path.name}: {exc}]"
    if len(text) > limit:
        text = text[:limit] + f"\n… [truncated {len(text) - limit} characters]"
    return text
